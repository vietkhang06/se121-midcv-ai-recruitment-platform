import os
import io
import pytest
from pathlib import Path
from fastapi.testclient import TestClient

from app.main import app
from app.config import settings
from app.services.document_extraction_service import DocumentExtractionService
from app.services.text_quality_evaluator import TextQualityEvaluator
from app.services.llm_client import LLMClient

FIXTURES_DIR = Path(__file__).resolve().parent.parent.parent / "test-fixtures" / "synthetic"


@pytest.fixture
def extraction_service():
    return DocumentExtractionService()


@pytest.fixture
def evaluator():
    return TextQualityEvaluator()


@pytest.fixture
def no_llm_guard(monkeypatch):
    """Proves strictly that NO LLM is ever called during document extraction."""
    def _prohibited_call(*args, **kwargs):
        raise AssertionError("CRITICAL VIOLATION: An LLM was invoked during document extraction!")

    monkeypatch.setattr(LLMClient, "generate_json", _prohibited_call)
    monkeypatch.setattr(LLMClient, "chat", _prohibited_call, raising=False)
    yield


class TestIsolatedDocumentExtraction:
    def test_single_column_pdf_extraction(self, extraction_service, no_llm_guard):
        file_path = str(FIXTURES_DIR / "01_single_column_tech_cv.pdf")
        assert os.path.exists(file_path)

        res = extraction_service.extract_from_file_path(file_path)
        assert res.success is True
        assert res.document.fileType == "PDF"
        assert res.document.extractionMethod == "PDF_TEXT"
        assert res.document.pageCount >= 1
        assert res.document.ocrUsed is False
        assert res.document.characterCount > 100
        assert res.document.wordCount > 20
        assert res.document.qualityScore > 0.8
        # Content assertions
        assert "NGUYEN TUAN ANH" in res.rawText
        assert "tuananh.nguyen@example.com" in res.rawText
        assert "Spring Boot" in res.rawText
        assert "FinTech Solutions Vietnam" in res.rawText

    def test_two_column_pdf_extraction(self, extraction_service, no_llm_guard):
        file_path = str(FIXTURES_DIR / "02_two_column_cv.pdf")
        assert os.path.exists(file_path)

        res = extraction_service.extract_from_file_path(file_path)
        assert res.success is True
        assert res.document.fileType == "PDF"
        assert res.document.characterCount > 100
        # Content assertions
        assert "LE MINH QUAN" in res.rawText
        assert "quan.le@example.com" in res.rawText
        assert "FastAPI" in res.rawText
        assert "NextGen Tech Hanoi" in res.rawText

    def test_docx_table_extraction(self, extraction_service, no_llm_guard):
        file_path = str(FIXTURES_DIR / "05_structured_table_cv.docx")
        assert os.path.exists(file_path)

        res = extraction_service.extract_from_file_path(file_path)
        assert res.success is True
        assert res.document.fileType == "DOCX"
        assert res.document.extractionMethod == "DOCX_PARSER"
        assert res.document.characterCount > 200
        assert res.document.ocrUsed is False
        # Content assertions
        assert "PHAM DUC THANG" in res.rawText
        assert "thang.pham@example.com" in res.rawText
        assert "Saigon Tech Labs" in res.rawText
        assert "InnoSoft Vietnam" in res.rawText
        assert "React" in res.rawText

    def test_image_png_extraction(self, extraction_service, no_llm_guard):
        file_path = str(FIXTURES_DIR / "04_cv_image.png")
        assert os.path.exists(file_path)

        res = extraction_service.extract_from_file_path(file_path)
        assert res.document.fileType == "IMAGE"
        # If Tesseract is installed, success=True, ocrUsed=True
        # If Tesseract not installed, handles clean error without crashing
        if res.success:
            assert res.document.ocrUsed is True
            assert res.document.extractionMethod == "IMAGE_OCR"
            assert "TRAN THI MAI" in res.rawText or res.document.characterCount > 0
        else:
            assert res.error_code in ["TESSERACT_NOT_FOUND", "OCR_FAILED"]

    def test_empty_file_fails_cleanly(self, extraction_service, no_llm_guard):
        res = extraction_service.extract_from_bytes(b"", file_name="empty.pdf")
        assert res.success is False
        assert res.error_code == "FILE_EMPTY"
        assert res.document.characterCount == 0

    def test_corrupted_file_fails_cleanly(self, extraction_service, no_llm_guard):
        corrupt_bytes = b"%PDF-1.4\nCorrupted content invalid trailer %%EOF"
        res = extraction_service.extract_from_bytes(corrupt_bytes, file_name="corrupt.pdf")
        assert res.success is False
        assert res.error_code in ["PDF_EXTRACTION_FAILED", "PDF_NO_PAGES"]

    def test_fake_extension_detected(self, extraction_service, no_llm_guard):
        # Plain text file disguised as a PDF
        fake_pdf = b"Just plain text pretending to be a PDF file."
        res = extraction_service.extract_from_bytes(fake_pdf, file_name="fake.pdf")
        # Magic bytes detect it is NOT valid PDF
        assert res.success is False
        assert res.error_code in ["UNSUPPORTED_FILE_TYPE", "PDF_EXTRACTION_FAILED"]

    def test_oversized_file_rejected(self, extraction_service, no_llm_guard):
        oversized = b"0" * (11 * 1024 * 1024)  # 11 MB > 10 MB limit
        res = extraction_service.extract_from_bytes(oversized, file_name="huge.pdf")
        assert res.success is False
        assert res.error_code == "FILE_TOO_LARGE"

    def test_unsupported_file_type_rejected(self, extraction_service, no_llm_guard):
        exe_file = b"MZ\x90\x00\x03\x00\x00\x00"
        res = extraction_service.extract_from_bytes(exe_file, file_name="malicious.exe")
        assert res.success is False
        assert res.error_code == "UNSUPPORTED_FILE_TYPE"

    def test_vietnamese_accented_cv_extraction(self, extraction_service, no_llm_guard):
        import docx
        doc = docx.Document()
        doc.add_heading("HỒ SƠ ỨNG VIÊN - KỸ SƯ PHẦN MỀM", 0)
        doc.add_paragraph("Họ và tên: NGUYỄN VĂN AN")
        doc.add_paragraph("Email: vanan.nguyen@example.com")
        doc.add_paragraph("Địa chỉ: Quận Cầu Giấy, Hà Nội, Việt Nam")
        doc.add_paragraph("Kỹ năng: Lập trình Java, Kiến trúc Microservices, Cơ sở dữ liệu PostgreSQL")
        bio = io.BytesIO()
        doc.save(bio)

        res = extraction_service.extract_from_bytes(bio.getvalue(), file_name="cv_tieng_viet.docx")
        assert res.success is True
        assert "NGUYỄN VĂN AN" in res.rawText
        assert "vanan.nguyen@example.com" in res.rawText
        assert "Hà Nội" in res.rawText
        assert "Cơ sở dữ liệu" in res.rawText
        assert res.document.qualityScore > 0.8

    def test_english_cv_extraction(self, extraction_service, no_llm_guard):
        import docx
        doc = docx.Document()
        doc.add_heading("CURRICULUM VITAE - CLOUD ARCHITECT", 0)
        doc.add_paragraph("Full Name: DAVID MILLER")
        doc.add_paragraph("Email: david.miller@example.com")
        doc.add_paragraph("Summary: Senior Cloud Architect with 10+ years experience in AWS, Kubernetes, Terraform.")
        bio = io.BytesIO()
        doc.save(bio)

        res = extraction_service.extract_from_bytes(bio.getvalue(), file_name="cv_english.docx")
        assert res.success is True
        assert "DAVID MILLER" in res.rawText
        assert "david.miller@example.com" in res.rawText
        assert "Kubernetes" in res.rawText
        assert res.document.qualityScore > 0.8


class TestDevExtractEndpoint:
    def test_dev_endpoint_disabled_by_default_returns_404(self, monkeypatch):
        monkeypatch.setattr(settings, "ENABLE_DEV_AI_TEST_ENDPOINTS", False)
        client = TestClient(app)
        res = client.post(
            "/internal/ai/dev/extract-text",
            files={"file": ("test.pdf", b"%PDF-1.4\nTest", "application/pdf")}
        )
        assert res.status_code == 404

    def test_dev_endpoint_enabled_extracts_successfully(self, monkeypatch, no_llm_guard):
        monkeypatch.setattr(settings, "ENABLE_DEV_AI_TEST_ENDPOINTS", True)
        client = TestClient(app)
        file_path = str(FIXTURES_DIR / "01_single_column_tech_cv.pdf")
        with open(file_path, "rb") as f:
            file_bytes = f.read()

        res = client.post(
            "/internal/ai/dev/extract-text",
            files={"file": ("01_single_column_tech_cv.pdf", file_bytes, "application/pdf")}
        )
        assert res.status_code == 200
        data = res.json()
        assert data["success"] is True
        assert data["document"]["fileType"] == "PDF"
        assert data["document"]["extractionMethod"] == "PDF_TEXT"
        assert data["document"]["characterCount"] > 100
        assert "NGUYEN TUAN ANH" in data["rawText"]
        assert "Spring Boot" in data["rawText"]
