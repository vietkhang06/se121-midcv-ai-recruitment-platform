import os
import io
import docx
from PIL import Image, ImageDraw, ImageFont

def build_pdf_stream(stream_content: str) -> bytes:
    stream_bytes = stream_content.encode("latin1", errors="replace")
    pdf_content = (
        b"%PDF-1.4\n"
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n"
        b"2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj\n"
        b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >> endobj\n"
        b"4 0 obj << /Length " + str(len(stream_bytes)).encode("latin1") + b" >> stream\n"
        + stream_bytes + b"\nendstream\nendobj\n"
        b"5 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n"
        b"xref\n0 6\n"
        b"0000000000 65535 f \n"
        b"0000000009 00000 n \n"
        b"0000000058 00000 n \n"
        b"0000000115 00000 n \n"
        b"0000000244 00000 n \n"
        b"0000000350 00000 n \n"
        b"trailer << /Size 6 /Root 1 0 R >>\nstartxref\n430\n%%EOF\n"
    )
    return pdf_content

def generate_fixtures(out_dir: str):
    os.makedirs(out_dir, exist_ok=True)

    # 1. Single Column Text PDF
    single_col_stream = (
        "BT\n/F1 12 Tf\n"
        "50 720 Td\n(NGUYEN TUAN ANH) Tj\n"
        "0 -20 Td\n(Senior Java Backend Developer) Tj\n"
        "0 -20 Td\n(Email: tuananh.nguyen@example.com | Phone: +84901234567 | Da Nang, Vietnam) Tj\n"
        "0 -25 Td\n(SUMMARY) Tj\n"
        "0 -15 Td\n(Over 5 years of software engineering experience specializing in Spring Boot, PostgreSQL, Kafka.) Tj\n"
        "0 -25 Td\n(SKILLS) Tj\n"
        "0 -15 Td\n(Java, Spring Boot, PostgreSQL, Docker, Redis, Kubernetes) Tj\n"
        "0 -25 Td\n(EXPERIENCE) Tj\n"
        "0 -15 Td\n(Senior Software Engineer - FinTech Solutions Vietnam (2021 - 2024)) Tj\n"
        "0 -15 Td\n(Architected payment gateway serving 500k daily transactions with 99.99% uptime.) Tj\n"
        "0 -25 Td\n(EDUCATION) Tj\n"
        "0 -15 Td\n(Bachelor of Software Engineering - Da Nang University of Technology (2020)) Tj\n"
        "ET"
    )
    with open(os.path.join(out_dir, "01_single_column_tech_cv.pdf"), "wb") as f:
        f.write(build_pdf_stream(single_col_stream))

    # 2. Two-Column Layout PDF
    two_col_stream = (
        "BT\n/F1 12 Tf\n"
        "50 720 Td\n(LE MINH QUAN) Tj\n"
        "0 -20 Td\n(quan.le@example.com) Tj\n"
        "0 -15 Td\n(+84988776655) Tj\n"
        "0 -25 Td\n(TECHNICAL SKILLS) Tj\n"
        "0 -15 Td\n(Python, FastAPI, Docker, PyTorch) Tj\n"
        "0 -25 Td\n(EDUCATION) Tj\n"
        "0 -15 Td\n(VNU University of Engineering) Tj\n"
        "0 -15 Td\n(Graduated: 2021) Tj\n"
        "300 135 Td\n(SENIOR AI & BACKEND DEVELOPER) Tj\n"
        "0 -20 Td\n(PROFESSIONAL SUMMARY) Tj\n"
        "0 -15 Td\n(Dedicated AI Engineer with deep focus on NLP, OCR, and scalable LLM microservices.) Tj\n"
        "0 -25 Td\n(WORK EXPERIENCE) Tj\n"
        "0 -15 Td\n(AI Lead - NextGen Tech Hanoi (2021 - 2024)) Tj\n"
        "0 -15 Td\n(Deployed document processing pipeline reducing extraction latency by 45%.) Tj\n"
        "ET"
    )
    with open(os.path.join(out_dir, "02_two_column_cv.pdf"), "wb") as f:
        f.write(build_pdf_stream(two_col_stream))

    # 3. Scanned PDF (Pure Image in PDF container)
    img_cv = Image.new("RGB", (800, 1000), color="white")
    draw = ImageDraw.Draw(img_cv)
    draw.text((50, 50), "TRAN THI MAI", fill="black")
    draw.text((50, 90), "Marketing Specialist & Brand Strategist", fill="black")
    draw.text((50, 130), "Email: mai.tran@example.com | Phone: 0912345678", fill="black")
    draw.text((50, 180), "Skills: Content Marketing, SEO, Google Analytics, English Fluent", fill="black")
    draw.text((50, 230), "Experience: Marketing Manager at VietBrands (2020 - 2024)", fill="black")
    draw.text((50, 280), "Education: Foreign Trade University, Bachelor of Business (2019)", fill="black")

    # Save as PNG
    img_cv.save(os.path.join(out_dir, "04_cv_image.png"), "PNG")

    # Save as Scanned PDF
    img_cv.save(os.path.join(out_dir, "03_scanned_cv.pdf"), "PDF", resolution=150.0)

    # 4. DOCX Document with Tables & Paragraphs
    doc = docx.Document()
    doc.add_heading("PHAM DUC THANG - Full Stack Developer", level=1)
    p_contact = doc.add_paragraph()
    p_contact.add_run("Email: thang.pham@example.com | Phone: 0977112233 | Ho Chi Minh City\n")
    p_contact.add_run("LinkedIn: https://linkedin.com/in/thangpham-dev | GitHub: https://github.com/thangpham-dev")

    doc.add_heading("Summary", level=2)
    doc.add_paragraph("Full stack engineer with expertise in React, Next.js, Node.js, and cloud deployments.")

    doc.add_heading("Skills", level=2)
    doc.add_paragraph("React, TypeScript, Next.js, Node.js, PostgreSQL, Docker, TailwindCSS")

    doc.add_heading("Work Experience", level=2)
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Company & Role"
    hdr_cells[1].text = "Timeline"
    hdr_cells[2].text = "Key Responsibilities"

    row_cells = table.add_row().cells
    row_cells[0].text = "Saigon Tech Labs\nSenior Frontend Engineer"
    row_cells[1].text = "2022 - 2024"
    row_cells[2].text = "Led development of core enterprise dashboard using Next.js and TypeScript."

    row_cells2 = table.add_row().cells
    row_cells2[0].text = "InnoSoft Vietnam\nFrontend Developer"
    row_cells2[1].text = "2020 - 2022"
    row_cells2[2].text = "Built responsive consumer web apps with React and state management."

    doc.add_heading("Education", level=2)
    doc.add_paragraph("Ho Chi Minh City University of Technology (HCMUT) - B.S. in Computer Science (2020)")

    doc.save(os.path.join(out_dir, "05_structured_table_cv.docx"))

    # 5. Legacy Word DOC file dummy with DOC magic bytes
    # Starts with OLE Compound File header "\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1"
    ole_header = b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1" + b"\x00" * 504
    legacy_text = b"HOANG VAN NAM - Project Manager\nEmail: nam.hoang@example.com\nSkills: Agile, Scrum, Jira"
    with open(os.path.join(out_dir, "06_legacy_cv.doc"), "wb") as f:
        f.write(ole_header + legacy_text)

    print(f"Generated 6 synthetic fixtures in {out_dir}")

if __name__ == "__main__":
    generate_fixtures(os.path.abspath("test-fixtures/synthetic"))
